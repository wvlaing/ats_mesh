from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ats_mesh.models import Job

COLUMNS: list[str] = ["ats", "company_name", "title", "posted", "locations", "url"]
WIDTHS: list[int] = [12, 25, 70, 14, 40, 50]
FONT = "Arial"


def posted_on(job: Job) -> str:
    """No date from the ATS: use the day we first saw it."""
    return job.posted or job.first_seen[:10]


def cell_value(job: Job, name: str) -> str | None:
    if name == "posted":
        return posted_on(job)
    value: Any | None = getattr(
        job, name, None
    )  # your own extra columns come back empty
    return "; ".join(value) if isinstance(value, list) else value


def open_sheet(path: Path) -> tuple[Workbook, Worksheet]:
    if path.exists():
        wb: Workbook = load_workbook(filename=path)
        ws = wb["Jobs"] if "Jobs" in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        assert isinstance(ws, Worksheet)
        ws.title = "Jobs"
        for col, (name, width) in enumerate(
            iterable=zip(COLUMNS, WIDTHS, strict=True), start=1
        ):
            head: Cell = ws.cell(row=1, column=col, value=name)
            head.font = Font(name=FONT, bold=True)
            head.fill = PatternFill("solid", start_color="DDE4EE")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "D2"  # keep ats, company and title visible
    assert isinstance(ws, Worksheet)
    return wb, ws


def sort_by_posted(ws: Worksheet, posted_col: int) -> None:
    """Reorder data rows by the posted column, most recent first.
    Rows are moved as a whole (all columns, including your own), the sort is
    stable, and rows with no date go to the bottom."""
    ncols: int = ws.max_column
    rows = [list(r) for r in ws.iter_rows(min_row=2, max_col=ncols, values_only=True)]
    rows.sort(key=lambda r: str(r[posted_col - 1] or ""), reverse=True)
    for offset, values in enumerate(rows):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=offset + 2, column=col)
            cell.hyperlink = None  # style_row re-adds it for the new row's url
            cell.value = value  # pyright: ignore[reportAttributeAccessIssue]


def style_row(ws: Worksheet, row: int, url_col: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="top")
    link = ws.cell(row=row, column=url_col)
    if isinstance(link, Cell) and str(object=link.value).startswith("http"):
        link.hyperlink = str(object=link.value)
        link.font = Font(name=FONT, color="0563C1", underline="single")


def sync(path: Path, added: list[Job], removed: list[Job]) -> None:
    """Append `added` jobs to the sheet and delete `removed` ones (the file is created
    if it doesn't exist), then sort every row by posted, most recent first.
    Everything else in the sheet (your own columns, notes) is left alone, which is
    why this isn't just a re-export.
    Rows are matched on the company_name and url columns."""
    wb, ws = open_sheet(path)

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    if "company_name" not in header or "url" not in header:
        raise SystemExit(f"{path} needs 'company_name' and 'url' columns in row 1")
    company_col: int = header.index("company_name") + 1
    url_col: int = header.index("url") + 1

    def key(row: int) -> tuple[str, str]:
        return (
            str(ws.cell(row=row, column=company_col).value),
            str(ws.cell(row=row, column=url_col).value),
        )

    gone: set[tuple[str, str]] = {(j.company_name, str(object=j.url)) for j in removed}
    deleted = 0
    for row in range(ws.max_row, 1, -1):  # bottom-up so deleting doesn't shift the rest
        if key(row) in gone:
            ws.delete_rows(row)
            deleted += 1

    have: set[tuple[str, str]] = {key(r) for r in range(2, ws.max_row + 1)}
    appended = 0
    for job in sorted(added, key=lambda j: (posted_on(j), j.first_seen), reverse=True):
        if (job.company_name, str(object=job.url)) not in have:
            ws.append([cell_value(job, str(name)) for name in header])
            appended += 1

    if "posted" in header:
        sort_by_posted(ws, header.index("posted") + 1)

    for row in range(2, ws.max_row + 1):
        style_row(ws, row, url_col, ws.max_column)
    ws.auto_filter.ref = ws.dimensions
    wb.save(filename=path)
    print(f"Excel: +{appended} / -{deleted} rows in {path.name}")
