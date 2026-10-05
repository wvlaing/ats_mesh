# ~/src/ats_mesh/helpers/export.py

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

COLUMNS = ["ats", "company_name", "title", "posted", "locations", "url"]
WIDTHS = [16, 22, 50, 14, 40, 50]
FONT = "Arial"


def cell_value(job: dict, name: str):
    value = job.get(name)
    if name == "posted":  # no date from the ATS: use the day we first saw it
        return value or job["first_seen"][:10]
    if isinstance(value, list):
        return "; ".join(value)
    return value


def new_sheet() -> tuple[Workbook, Worksheet]:
    wb = Workbook()
    ws = wb.active
    assert isinstance(ws, Worksheet)
    ws.title = "Jobs"
    for i, (name, width) in enumerate(zip(COLUMNS, WIDTHS), start=1):
        head = ws.cell(row=1, column=i, value=name)
        head.font = Font(name=FONT, bold=True)
        head.fill = PatternFill("solid", start_color="DDE4EE")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"  # keep ats, company and title visible
    return wb, ws


def open_sheet(path: Path) -> tuple[Workbook, Worksheet]:
    if not path.exists():
        return new_sheet()
    wb = load_workbook(path)
    ws = wb["Jobs"] if "Jobs" in wb.sheetnames else wb.active
    assert isinstance(ws, Worksheet)
    return wb, ws


def style_row(ws: Worksheet, row: int, url_col: int) -> None:
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT)
        cell.alignment = Alignment(vertical="top")
    link = ws.cell(row=row, column=url_col)
    if isinstance(link, Cell) and link.value and str(link.value).startswith("http"):
        link.hyperlink = str(link.value)
        link.font = Font(name=FONT, color="0563C1", underline="single")


def sync(path: Path, added: list[dict], removed: list[dict]) -> None:
    """Append `added` jobs to the sheet and delete `removed` ones (the file is created
    if it doesn't exist). Everything else in the sheet (your own columns, notes,
    sorting) is left alone, which is why this isn't just a re-export.
    Rows are matched on the company_name and url columns."""
    wb, ws = open_sheet(path)

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    if "company_name" not in header or "url" not in header:
        raise SystemExit(f"{path} needs 'company_name' and 'url' columns in row 1")
    company_col = header.index("company_name") + 1
    url_col = header.index("url") + 1

    def key(row: int) -> tuple[str, str]:
        return (
            str(ws.cell(row=row, column=company_col).value),
            str(ws.cell(row=row, column=url_col).value),
        )

    gone = {(str(j["company_name"]), str(j["url"])) for j in removed}
    deleted = 0
    for row in range(ws.max_row, 1, -1):  # bottom-up so deleting doesn't shift the rest
        if key(row) in gone:
            ws.delete_rows(row)
            deleted += 1

    have = {key(r) for r in range(2, ws.max_row + 1)}
    newest_first = sorted(
        added, key=lambda j: (cell_value(j, "posted"), j["first_seen"]), reverse=True
    )
    appended = 0
    for job in newest_first:
        if (str(job["company_name"]), str(job["url"])) not in have:
            ws.append([cell_value(job, str(name)) for name in header])
            appended += 1

    for row in range(2, ws.max_row + 1):  # delete_rows loses hyperlinks, so redo all
        style_row(ws, row, url_col)
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    print(f"Excel: +{appended} / -{deleted} rows in {path.name}")
